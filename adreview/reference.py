"""점검 기준 문서(HWP / HWPX / DOCX / PDF / TXT) → 조항 단위 색인 → 관련 조항 검색.

LLM 없이 동작한다. 검색은 한글 글자 2-gram 기반 BM25 (순수 파이썬, 추가 패키지 불필요).
역할은 "판정"이 아니라 "광고 문구와 관련 있는 조항을 담당자에게 보여주기" 이다.

HWP(5.x 바이너리) 처리 순서
  1) olefile 로 BodyText/Section* 을 직접 파싱 (이 파일의 parse_hwp_records)
  2) 실패 시 `hwp5txt`(pyhwp) 가 설치돼 있으면 그것을 호출
  3) 그래도 안 되면 PrvText(미리보기 텍스트, 앞부분 약 1KB 수준)만 사용하고 경고
  ※ 배포용(암호화) HWP 는 지원하지 않는다. 한글에서 일반 문서로 다시 저장하거나 DOCX/PDF 로 변환해 넣을 것.
"""
from __future__ import annotations

import difflib
import json
import math
import re
import shutil
import struct
import subprocess
import zlib
from collections import Counter
from pathlib import Path
from typing import Optional

_SRC_PRIORITY = {".docx": 0, ".hwpx": 1, ".hwp": 2, ".md": 3, ".txt": 3, ".pdf": 4}
REF_EXTS = {".hwp", ".hwpx", ".docx", ".pdf", ".txt", ".md"}

# ----------------------------------------------------------------------------
# HWP
# ----------------------------------------------------------------------------
_HWPTAG_PARA_TEXT = 0x10 + 51

# 제어 문자 중 1 WCHAR 만 차지하는 것들, 그 외 (<32) 는 8 WCHAR(16바이트) 확장 제어
_CHAR_CTRL = {0, 10, 13, 24, 25, 26, 27, 28, 29, 30, 31}


def _decode_para_text(data: bytes) -> str:
    out = []
    i, n = 0, len(data)
    while i + 1 < n:
        code = struct.unpack_from("<H", data, i)[0]
        if code >= 32:
            out.append(chr(code))
            i += 2
        elif code in _CHAR_CTRL:
            if code in (10, 13):
                out.append("\n")
            elif code == 24:
                out.append("-")
            elif code in (30, 31):
                out.append(" ")
            i += 2
        else:
            if code == 9:
                out.append("\t")
            i += 16  # 확장/인라인 제어: 코드 2바이트 + 정보 14바이트
    return "".join(out)


def parse_hwp_records(stream: bytes) -> list:
    """BodyText 섹션 스트림(압축 해제 후)에서 문단 텍스트 목록을 뽑는다."""
    paras = []
    pos, n = 0, len(stream)
    while pos + 4 <= n:
        header = struct.unpack_from("<I", stream, pos)[0]
        pos += 4
        tag = header & 0x3FF
        size = (header >> 20) & 0xFFF
        if size == 0xFFF:
            if pos + 4 > n:
                break
            size = struct.unpack_from("<I", stream, pos)[0]
            pos += 4
        if pos + size > n:
            break
        if tag == _HWPTAG_PARA_TEXT:
            text = _decode_para_text(stream[pos:pos + size]).strip()
            if text:
                paras.append(text)
        pos += size
    return paras


def read_hwp(path: str) -> tuple:
    """(문단 목록, 경고 목록)"""
    import olefile

    warnings = []
    try:
        ole = olefile.OleFileIO(path)
    except Exception as e:
        return [], [f"HWP 열기 실패(HWP 5.x 형식이 아닐 수 있음): {e}"]

    try:
        header = ole.openstream("FileHeader").read()
        flags = struct.unpack_from("<I", header, 36)[0]
        compressed, encrypted = bool(flags & 1), bool(flags & 2)
        if encrypted or (flags & 4):  # 암호화 또는 배포용 문서
            warnings.append("암호화/배포용 HWP 입니다. 일반 문서로 다시 저장하거나 DOCX/PDF 로 변환해 주세요.")
            return [], warnings

        sections = sorted(
            (e for e in ole.listdir() if len(e) == 2 and e[0] == "BodyText" and e[1].startswith("Section")),
            key=lambda e: int(re.sub(r"\D", "", e[1]) or 0),
        )
        paras = []
        for entry in sections:
            raw = ole.openstream(entry).read()
            if compressed:
                raw = zlib.decompressobj(-15).decompress(raw, 256 * 1024 * 1024)   # 압축 폭탄 방지(상한 256MB)
            paras.extend(parse_hwp_records(raw))
        if paras:
            return paras, warnings
        warnings.append("HWP 본문에서 텍스트를 찾지 못했습니다.")
    except Exception as e:
        warnings.append(f"HWP 직접 파싱 실패: {e}")
    finally:
        ole.close()

    # 2) hwp5txt
    exe = shutil.which("hwp5txt")
    if exe:
        try:
            r = subprocess.run([exe, path], capture_output=True, timeout=120)
            txt = r.stdout.decode("utf-8", errors="ignore")
            paras = [l.strip() for l in txt.splitlines() if l.strip()]
            if paras:
                warnings.append("hwp5txt 로 대체 추출함")
                return paras, warnings
        except Exception as e:
            warnings.append(f"hwp5txt 실패: {e}")

    # 3) PrvText
    try:
        ole = olefile.OleFileIO(path)
        if ole.exists("PrvText"):
            txt = ole.openstream("PrvText").read().decode("utf-16le", errors="ignore")
            paras = [l.strip() for l in txt.splitlines() if l.strip()]
            warnings.append("미리보기(PrvText)만 추출했습니다. 본문 전체가 아니므로 DOCX/PDF 변환본 사용을 권장합니다.")
            ole.close()
            return paras, warnings
        ole.close()
    except Exception:
        pass
    return [], warnings


def read_hwpx(path: str) -> tuple:
    import zipfile
    import xml.etree.ElementTree as ET

    paras = []
    with zipfile.ZipFile(path) as z:
        names = sorted(n for n in z.namelist() if re.match(r"Contents/section\d+\.xml", n))
        for name in names:
            root = ET.fromstring(z.read(name))
            for p in root.iter():
                if p.tag.endswith("}p"):
                    text = "".join(t.text or "" for t in p.iter() if t.tag.endswith("}t")).strip()
                    if text:
                        paras.append(text)
    return paras, []


# ----------------------------------------------------------------------------
# DOCX / PDF / TXT
# ----------------------------------------------------------------------------
def read_docx(path: str) -> tuple:
    import docx

    d = docx.Document(path)
    paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = []
            for c in row.cells:
                txt = c.text.strip().replace("\n", " ")
                if txt and (not cells or cells[-1] != txt):  # 병합 셀 중복 제거
                    cells.append(txt)
            if cells:
                paras.append(" | ".join(cells))
    return paras, []


def read_pdf(path: str) -> tuple:
    import pymupdf

    paras, warnings = [], []
    doc = pymupdf.open(path)
    total = 0
    for pg in doc:
        txt = pg.get_text("text")
        total += len(txt.strip())
        paras.extend(l.strip() for l in txt.splitlines() if l.strip())
    doc.close()
    if total == 0:
        warnings.append("텍스트 레이어가 없는 스캔 PDF 입니다. OCR 이 필요합니다 (python -m adreview ocr 로 페이지 이미지를 먼저 변환).")
    return paras, warnings


def read_text(path: str) -> tuple:
    for enc in ("utf-8", "cp949", "utf-16"):
        try:
            txt = Path(path).read_text(encoding=enc)
            return [l.strip() for l in txt.splitlines() if l.strip()], []
        except Exception:
            continue
    return [], ["텍스트 인코딩을 판별하지 못했습니다."]


def read_reference_file(path: str) -> tuple:
    ext = Path(path).suffix.lower()
    if ext == ".hwp":
        return read_hwp(path)
    if ext == ".hwpx":
        return read_hwpx(path)
    if ext == ".docx":
        return read_docx(path)
    if ext == ".pdf":
        return read_pdf(path)
    if ext in (".txt", ".md"):
        return read_text(path)
    return [], [f"지원하지 않는 형식: {ext}"]


# ----------------------------------------------------------------------------
# 조항 분할
# ----------------------------------------------------------------------------
_ARTICLE = re.compile(r"^\s*(제\s*\d+\s*조(?:의\s*\d+)?)\s*[\(（]?([^)）\n]*)[\)）]?")
_CHAPTER = re.compile(r"^\s*(제\s*\d+\s*[장절편])\s*(.*)")
_MAX_CLAUSE_CHARS = 1200


def split_clauses(paras: list, source: str) -> list:
    """'제N조(제목)' 단위로 묶는다. 조 번호가 없는 문서는 문단 단위로 적당히 묶는다."""
    clauses = []
    cur = None
    chapter = ""
    saw_article = any(_ARTICLE.match(p) for p in paras)

    def flush():
        nonlocal cur
        if cur and cur["text"].strip():
            clauses.append(cur)
        cur = None

    for p in paras:
        m_ch = _CHAPTER.match(p)
        if m_ch and len(p) < 60:
            chapter = p.strip()
            continue
        m = _ARTICLE.match(p)
        if m:
            flush()
            title = f"{m.group(1).replace(' ', '')}({m.group(2).strip()})" if m.group(2).strip() else m.group(1).replace(" ", "")
            cur = {"source": source, "chapter": chapter, "title": title, "text": p}
        elif saw_article:
            if cur is None:
                cur = {"source": source, "chapter": chapter, "title": "(조 이전 내용)", "text": p}
            else:
                cur["text"] += "\n" + p
        else:
            if cur is None or len(cur["text"]) > 400:
                flush()
                cur = {"source": source, "chapter": chapter, "title": f"문단 {len(clauses) + 1}", "text": p}
            else:
                cur["text"] += "\n" + p

    flush()

    # 너무 긴 조항은 분할
    out = []
    for c in clauses:
        if len(c["text"]) <= _MAX_CLAUSE_CHARS:
            out.append(c)
            continue
        lines, buf, part = c["text"].split("\n"), "", 1
        for ln in lines:
            if len(buf) + len(ln) > _MAX_CLAUSE_CHARS and buf:
                out.append({**c, "title": f"{c['title']} [{part}]", "text": buf})
                buf, part = "", part + 1
            buf += ("\n" if buf else "") + ln
        if buf:
            out.append({**c, "title": f"{c['title']} [{part}]" if part > 1 else c["title"], "text": buf})
    return out


# ----------------------------------------------------------------------------
# 검색 (2-gram BM25)
# ----------------------------------------------------------------------------
def _tokens(text: str) -> list:
    t = re.sub(r"[^0-9A-Za-z가-힣]", "", text)
    if len(t) < 2:
        return [t] if t else []
    return [t[i:i + 2] for i in range(len(t) - 1)]


class ReferenceIndex:
    def __init__(self, clauses: Optional[list] = None, warnings: Optional[list] = None):
        self.clauses = clauses or []
        self.warnings = warnings or []
        self._build()

    def _build(self, k1: float = 1.5, b: float = 0.75):
        self._tf = [Counter(_tokens(c["title"] + " " + c["text"])) for c in self.clauses]
        self._len = [sum(tf.values()) for tf in self._tf]
        self._avg = (sum(self._len) / len(self._len)) if self._len else 1.0
        df = Counter()
        for tf in self._tf:
            df.update(tf.keys())
        n = len(self.clauses)
        self._idf = {t: math.log(1 + (n - d + 0.5) / (d + 0.5)) for t, d in df.items()}
        self._k1, self._b = k1, b

    def search(self, query: str, k: int = 3, min_score: float = 1.0) -> list:
        q = Counter(_tokens(query))
        if not q or not self.clauses:
            return []
        scores = []
        for i, tf in enumerate(self._tf):
            s = 0.0
            for t in q:
                f = tf.get(t)
                if not f:
                    continue
                denom = f + self._k1 * (1 - self._b + self._b * self._len[i] / self._avg)
                s += self._idf.get(t, 0) * f * (self._k1 + 1) / denom
            scores.append((s, i))
        scores.sort(reverse=True)
        res, seen_norm = [], []
        for s, i in scores:
            if s < min_score or len(res) >= k:
                break
            c = self.clauses[i]
            norm = re.sub(r"[^0-9A-Za-z가-힣]", "", c["text"])[:200]
            dup = next((r for r, n in zip(res, seen_norm)
                        if difflib.SequenceMatcher(None, n, norm).ratio() >= 0.85), None)
            if dup is not None:      # 같은 규정이 HWP/DOCX/PDF 로 중복 보관된 경우 하나만 보여주고 출처만 합침
                # 표시용 본문은 서식·띄어쓰기가 잘 보존되는 형식(docx > hwpx > hwp > pdf)을 우선
                if _SRC_PRIORITY.get(Path(c["source"]).suffix.lower(), 9) < _SRC_PRIORITY.get(Path(dup["source"]).suffix.lower(), 9):
                    dup["also_in"].append(dup["source"])
                    dup.update(source=c["source"], title=c["title"], chapter=c["chapter"], text=c["text"])
                elif c["source"] not in dup["also_in"] and c["source"] != dup["source"]:
                    dup["also_in"].append(c["source"])
                continue
            seen_norm.append(norm)
            res.append({"source": c["source"], "title": c["title"], "chapter": c["chapter"],
                        "text": c["text"], "score": round(s, 2), "also_in": []})
        return res

    # ---- 저장/불러오기 ----
    def save(self, path: str):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        Path(path).write_text(json.dumps({"clauses": self.clauses, "warnings": self.warnings},
                                         ensure_ascii=False, indent=1), encoding="utf-8")

    @classmethod
    def load(cls, path: str) -> "ReferenceIndex":
        d = json.loads(Path(path).read_text(encoding="utf-8"))
        return cls(d["clauses"], d.get("warnings", []))


def build_reference_index(folder: str) -> ReferenceIndex:
    clauses, warnings = [], []
    files = sorted(p for p in Path(folder).rglob("*") if p.suffix.lower() in REF_EXTS and p.is_file())
    if not files:
        warnings.append(f"{folder} 에 기준 문서(hwp/hwpx/docx/pdf/txt)가 없습니다.")
    for f in files:
        paras, w = read_reference_file(str(f))
        warnings.extend(f"[{f.name}] {x}" for x in w)
        if not paras:
            warnings.append(f"[{f.name}] 추출된 텍스트가 없어 색인에서 제외됨")
            continue
        clauses.extend(split_clauses(paras, f.name))
    return ReferenceIndex(clauses, warnings)
