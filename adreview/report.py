"""결과 출력: 위치 표시 이미지 + HTML 리포트 + JSON + 엑셀."""
from __future__ import annotations

import html
import json
import os
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

from .checks import FileResult

SEV_COLOR = {"high": (220, 38, 38), "medium": (234, 138, 0), "low": (37, 99, 235), "info": (120, 120, 120)}
SEV_KO = {"high": "높음", "medium": "중간", "low": "낮음", "info": "참고"}

_FONT_CANDIDATES = [
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf",
    "C:/Windows/Fonts/malgunbd.ttf",
    "/System/Library/Fonts/AppleSDGothicNeo.ttc",
]


def _font(size: int):
    for p in _FONT_CANDIDATES:
        if os.path.exists(p):
            try:
                return ImageFont.truetype(p, size)
            except Exception:
                continue
    return ImageFont.load_default()


def annotate_images(result: FileResult, out_dir: str) -> dict:
    """finding 의 bbox 를 이미지에 그려서 저장. {(page_id, image_key): 상대경로} 반환."""
    os.makedirs(out_dir, exist_ok=True)
    mapping = {}
    for page in result.pages:
        by_img: dict = {}
        for idx, f in enumerate(x for x in result.findings if x.page == page.page_id):
            if f.bbox and f.image in page.images and f.status == "review":
                by_img.setdefault(f.image, []).append(f)
        for key, path in page.images.items():
            try:
                im = Image.open(path).convert("RGB")
            except Exception:
                continue
            draw = ImageDraw.Draw(im)
            fs = max(14, im.width // 70)
            font = _font(fs)
            # 거의 같은 위치의 finding 은 라벨을 하나로 합친다 (예: 한 줄에 BAN-001, BAN-002 가 같이 걸린 경우)
            clusters: dict = {}
            for f in by_img.get(key, []):
                ck = tuple(round(v / 12) for v in f.bbox)
                clusters.setdefault(ck, []).append(f)
            order = ["high", "medium", "low", "info"]
            for fl in clusters.values():
                top = min(fl, key=lambda x: order.index(x.severity) if x.severity in order else 9)
                x0 = min(f.bbox[0] for f in fl)
                y0 = min(f.bbox[1] for f in fl)
                x1 = max(f.bbox[2] for f in fl)
                y1 = max(f.bbox[3] for f in fl)
                pad = max(3, im.width // 400)
                col = SEV_COLOR.get(top.severity, (0, 0, 0))
                draw.rectangle([x0 - pad, y0 - pad, x1 + pad, y1 + pad], outline=col, width=max(3, im.width // 450))
                tag = " · ".join(dict.fromkeys(f.rule_id for f in fl))
                tw = draw.textlength(tag, font=font)
                ty = max(0, y0 - pad - fs - 6)
                draw.rectangle([x0 - pad, ty, x0 - pad + tw + 8, ty + fs + 6], fill=col)
                draw.text((x0 - pad + 4, ty + 2), tag, fill=(255, 255, 255), font=font)
            safe = f"{Path(page.file).stem}_{page.label}_{key}".replace(" ", "_").replace("/", "_")
            out = os.path.join(out_dir, safe + ".png")
            im.save(out)
            mapping[(page.page_id, key)] = os.path.basename(out)
    return mapping


def _esc(s) -> str:
    return html.escape(str(s if s is not None else ""))


_CSS = """
body{font-family:'Noto Sans CJK KR','Malgun Gothic','Apple SD Gothic Neo',sans-serif;margin:24px;color:#222;max-width:1200px}
h1{font-size:22px} h2{font-size:18px;border-bottom:2px solid #333;padding-bottom:4px;margin-top:36px}
h3{font-size:15px;margin:22px 0 6px}
.note{background:#fff7e0;border:1px solid #f0d890;padding:10px 14px;border-radius:6px;font-size:13px}
.sum span{display:inline-block;padding:3px 10px;border-radius:12px;color:#fff;font-size:13px;margin-right:6px}
.high{background:#dc2626}.medium{background:#ea8a00}.low{background:#2563eb}.info{background:#777}.ok{background:#16a34a}
table{border-collapse:collapse;width:100%;font-size:13px;margin:8px 0}
th,td{border:1px solid #ddd;padding:6px 8px;vertical-align:top;text-align:left}
th{background:#f3f3f3}
.rel{font-size:12px;color:#444;background:#f7f7f7;border-left:3px solid #bbb;padding:4px 8px;margin:4px 0;white-space:pre-wrap}
.imgs img{max-width:100%;border:1px solid #ccc;margin:6px 0}
details{margin:6px 0} summary{cursor:pointer;font-weight:600}
.small{color:#666;font-size:12px}
.mk{display:inline-block;min-width:26px;text-align:center;font-weight:700;border-radius:4px;padding:1px 6px;color:#fff}
.mk-o{background:#16a34a}.mk-x{background:#dc2626}.mk-t{background:#ea8a00}.mk-n{background:#999}.mk-q{background:#2563eb}
.chk td:nth-child(1){white-space:nowrap} .chk td:nth-child(4){text-align:center}
td:nth-child(1),td:nth-child(2){white-space:nowrap}
td:nth-child(3){min-width:180px}
"""


_MK_CLASS = {"○": "mk-o", "×": "mk-x", "△": "mk-t", "-": "mk-n", "?": "mk-q"}


def _checklist_html(rows: list) -> str:
    from .checklist import MARK_KO, summarize
    sm = summarize(rows)
    out = ["<h3>광고심의점검표 (변액보험 이외 보험) — 자동 판정 초안</h3>",
           "<p class='small'>○ 적합 · × 부적합 · △ 미흡 · - 해당사항 없음 · ? 자동판정 불가(담당자 확인). "
           "× 는 문구 누락/금지 표현 검출이며 OCR 누락일 수 있고, ○ 는 자동 패턴이 포착한 범위 안에서만 의미가 있습니다.</p>",
           "<div class='sum'>" + "".join(
               f"<span class='{_MK_CLASS[m]}'>{m} {MARK_KO[m]} {n}</span>" for m, n in sm.items()) + "</div>",
           "<table class='chk'><tr><th>No.</th><th>구분</th><th>점검 항목</th><th>확인</th><th>자동 판정 근거 / 비고</th><th>규정</th></tr>"]
    last = None
    for r in rows:
        grp = r.group if r.group != last else ""
        last = r.group
        ev = f"<div class='small'>{_esc(' · '.join(r.evidence[:4]))}</div>" if r.evidence else ""
        sample = f"<details><summary class='small'>예시 문구</summary><div class='rel'>{_esc(r.sample)}</div></details>" if r.sample else ""
        kind = "" if r.mode == "auto" else "<div class='small'>(담당자 확인 항목)</div>"
        out.append(f"<tr><td>{_esc(r.no)}</td><td>{_esc(grp)}</td><td>{_esc(r.title)}{sample}</td>"
                   f"<td><span class='mk {_MK_CLASS[r.mark]}'>{_esc(r.mark)}</span></td>"
                   f"<td>{_esc(r.detail)}{ev}{kind}</td><td class='small'>{_esc(r.basis)}</td></tr>")
    out.append("</table>")
    return "".join(out)


def write_html(results: list, out_dir: str, image_maps: dict, title: str = "광고물 1차 점검 리포트") -> str:
    parts = [f"<!doctype html><html lang='ko'><head><meta charset='utf-8'><title>{_esc(title)}</title><style>{_CSS}</style></head><body>",
             f"<h1>{_esc(title)}</h1>",
             "<div class='note'><b>이 리포트는 심의 담당자를 돕는 1차 점검 결과입니다.</b> "
             "‘의심’으로 표시된 항목은 위반 확정이 아니며, 누락 판정은 OCR 한계로 오탐일 수 있습니다. "
             "최종 판단은 반드시 담당자가 원본과 규정을 대조해 내려야 합니다.</div>"]
    for res in results:
        review = [f for f in res.findings if f.status == "review"]
        oks = [f for f in res.findings if f.status == "ok"]
        infos = [f for f in res.findings if f.status == "info"]
        cnt = {s: sum(1 for f in review if f.severity == s) for s in ("high", "medium", "low")}
        parts.append(f"<h2>{_esc(res.file)}</h2>")
        parts.append("<div class='sum'>"
                     f"<span class='high'>높음 {cnt['high']}</span><span class='medium'>중간 {cnt['medium']}</span>"
                     f"<span class='low'>낮음 {cnt['low']}</span><span class='ok'>확인됨 {len(oks)}</span></div>")
        meta = f"페이지 {len(res.pages)}개"
        if res.products:
            meta += f" · 식별 상품: {', '.join(res.products)}"
        if res.product_types:
            meta += f" · 상품군: {', '.join(res.product_types)}"
        parts.append(f"<p class='small'>{_esc(meta)}</p>")
        if res.skipped_rules:
            parts.append(f"<p class='small'>적용 제외 룰: {_esc('; '.join(res.skipped_rules))}</p>")

        if res.checklist:
            parts.append(_checklist_html(res.checklist))

        parts.append("<h3>검토 필요 항목</h3>")
        if not review:
            parts.append("<p>자동 점검에서 걸린 항목이 없습니다. (그렇다고 규정에 부합한다는 뜻은 아닙니다.)</p>")
        else:
            parts.append("<table><tr><th>중요도</th><th>룰</th><th>항목</th><th>위치</th><th>상세</th></tr>")
            for f in review:
                rel = ""
                if f.related:
                    rel = "<details><summary>관련 기준 조항 (검색 결과)</summary>" + "".join(
                        f"<div class='rel'><b>{_esc(r['source'])}{(' (동일 내용: ' + ', '.join(r['also_in']) + ')') if r.get('also_in') else ''} · {_esc(r['title'])}</b> (유사도 {r['score']})\n{_esc(r['text'][:400])}{'…' if len(r['text']) > 400 else ''}</div>"
                        for r in f.related) + "</details>"
                basis = f"<div class='small'>근거: {_esc(f.basis)}</div>" if f.basis else ""
                parts.append(
                    f"<tr><td><span class='{f.severity}' style='color:#fff;padding:2px 8px;border-radius:10px'>{SEV_KO.get(f.severity, f.severity)}</span></td>"
                    f"<td>{_esc(f.rule_id)}</td><td>{_esc(f.title)}</td><td>{_esc(f.page)}</td>"
                    f"<td>{_esc(f.detail)}{basis}{rel}</td></tr>")
            parts.append("</table>")

        # 위치 표시 이미지
        shown = False
        for page in res.pages:
            for key in page.images:
                fn = image_maps.get((page.page_id, key))
                has = any(f.bbox and f.image == key and f.page == page.page_id and f.status == "review" for f in res.findings)
                if fn and has:
                    if not shown:
                        parts.append("<h3>위치 표시 (박스 = 검토 필요 지점)</h3><div class='imgs'>")
                        shown = True
                    parts.append(f"<div class='small'>{_esc(page.page_id)} · {_esc(key)}</div><img src='images/{_esc(fn)}'>")
        if shown:
            parts.append("</div>")

        if oks:
            parts.append("<details><summary>확인된 항목 (필수 문구·수치 일치)</summary><table><tr><th>룰</th><th>항목</th><th>위치</th><th>상세</th></tr>")
            for f in oks:
                parts.append(f"<tr><td>{_esc(f.rule_id)}</td><td>{_esc(f.title)}</td><td>{_esc(f.page)}</td><td>{_esc(f.detail)}</td></tr>")
            parts.append("</table></details>")
        if infos:
            parts.append("<details><summary>참고·처리 경고</summary><ul>" +
                         "".join(f"<li>{_esc(f.page)} — {_esc(f.title)}: {_esc(f.detail)}</li>" for f in infos) + "</ul></details>")

        parts.append("<details><summary>추출된 텍스트 (OCR/네이티브)</summary>")
        for page in res.pages:
            parts.append(f"<h4>{_esc(page.page_id)}</h4><pre class='small' style='white-space:pre-wrap'>")
            seen_txt = set()
            for b in page.blocks:
                k = (b.source, b.image, "".join(b.text.split()))
                if k in seen_txt:
                    continue
                seen_txt.add(k)
                tag = "N" if b.source == "native" else f"O:{b.image or ''}/{b.variant or '-'}"
                parts.append(_esc(f"[{tag} {b.conf:.0f}] {b.text}\n"))
            parts.append("</pre>")
        parts.append("</details>")
    parts.append("</body></html>")
    path = os.path.join(out_dir, "report.html")
    Path(path).write_text("".join(parts), encoding="utf-8")
    return path


def write_json(results: list, out_dir: str) -> str:
    data = []
    for r in results:
        data.append({
            "file": r.file, "products": r.products, "product_types": r.product_types,
            "skipped_rules": r.skipped_rules,
            "checklist": [c.to_dict() for c in r.checklist],
            "findings": [f.to_dict() for f in r.findings],
            "pages": [{"page": p.page_id, "warnings": p.warnings,
                       "blocks": [b.to_dict() for b in p.blocks]} for p in r.pages],
        })
    path = os.path.join(out_dir, "result.json")
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    return path


def write_xlsx(results: list, out_dir: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "검토 항목"
    head = ["파일", "페이지", "상태", "중요도", "룰", "항목", "매칭 문구", "상세", "근거 규정", "관련 조항(상위1)", "담당자 판단", "비고"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="333333")
    fills = {"high": "FDE2E2", "medium": "FFF1D6", "low": "E3EDFF", "info": "EEEEEE"}
    for r in results:
        for f in r.findings:
            rel = f"{f.related[0]['source']} {f.related[0]['title']}: {f.related[0]['text'][:150]}" if f.related else ""
            status = {"review": "검토필요", "ok": "확인됨", "info": "참고"}.get(f.status, f.status)
            ws.append([r.file, f.page, status, SEV_KO.get(f.severity, f.severity), f.rule_id, f.title,
                       f.matched_text, f.detail, f.basis, rel, "", ""])
            if f.status == "review":
                for c in ws[ws.max_row]:
                    c.fill = PatternFill("solid", fgColor=fills.get(f.severity, "FFFFFF"))
    widths = [22, 18, 9, 8, 12, 36, 20, 60, 30, 50, 16, 20]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    if any(r.checklist for r in results):
        _checklist_sheet(wb, results)
    path = os.path.join(out_dir, "result.xlsx")
    wb.save(path)
    return path


def _checklist_sheet(wb, results: list) -> None:
    """협회 제출양식(No./점검항목/확인/비고/예시 문구)에 가깝게 점검표 시트를 만든다."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from .checklist import MARK_KO
    ws = wb.create_sheet("광고심의점검표")
    ws.append(["※ 적합 ○ / 부적합 × / 미흡 △ / 해당사항 없음 - / 자동판정 불가 ?  — 자동 판정은 초안이며 담당자가 확인·수정해야 합니다."])
    ws.append(["파일", "No.", "구분", "점검 항목", "확인(자동)", "판정 의미", "자동 판정 근거", "비고(규정)", "예시 문구", "담당자 확인", "담당자 의견"])
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="333333")
    fills = {"○": "DCFCE7", "×": "FDE2E2", "△": "FFF1D6", "?": "E3EDFF", "-": "EEEEEE"}
    for r in results:
        for c in r.checklist:
            ws.append([r.file, c.no, c.group, c.title, c.mark, MARK_KO[c.mark],
                       c.detail + ((" | " + " · ".join(c.evidence[:4])) if c.evidence else ""),
                       c.basis, c.sample, "", ""])
            ws.cell(ws.max_row, 5).fill = PatternFill("solid", fgColor=fills.get(c.mark, "FFFFFF"))
            ws.cell(ws.max_row, 5).font = Font(bold=True)
    for i, w in enumerate([22, 7, 12, 40, 10, 12, 70, 18, 50, 12, 24]):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A3"
