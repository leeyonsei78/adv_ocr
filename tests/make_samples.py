"""테스트용 샘플 광고물(이미지/PPTX)과 가상의 기준 문서(DOCX/PDF)를 만든다.
샘플 내용은 모두 가상이며, 일부러 위반 의심 요소를 심어 두었다."""
import os
import subprocess
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
S = ROOT / "samples"
R = ROOT / "reference_docs"
S.mkdir(exist_ok=True)
R.mkdir(exist_ok=True)

FONT = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
FONT_B = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"


def f(size, bold=False):
    return ImageFont.truetype(FONT_B if bold else FONT, size)


def banner_bad(path):
    """위반 요소를 일부러 넣은 암보험 배너 (밝은 배경 + 어두운 배경 영역 혼합)."""
    im = Image.new("RGB", (1400, 900), (255, 255, 255))
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, 1400, 330], fill=(20, 40, 90))
    d.text((60, 50), "든든암보험", font=f(54, True), fill=(255, 220, 80))
    d.text((60, 130), "대한민국 최고의 암보험! 무조건 100% 보장", font=f(62, True), fill=(255, 255, 255))
    d.text((60, 235), "타사보다 저렴한 보험료, 지금 바로 가입하세요", font=f(40), fill=(200, 220, 255))
    d.text((60, 380), "암 진단비 최대 7,000만원", font=f(56, True), fill=(200, 30, 30))
    d.text((60, 470), "소액암 진단비 1,000만원 · 가입연령 20~70세", font=f(40), fill=(30, 30, 30))
    d.text((60, 540), "20년납 20년만기 · 보장개시 후 90일 면책", font=f(40), fill=(30, 30, 30))
    d.text((60, 780), "※ 자세한 보장내용은 약관을 확인하시기 바랍니다.", font=f(15), fill=(120, 120, 120))
    d.text((60, 805), "※ 갱신형 상품으로 갱신 시 보험료가 인상될 수 있습니다.", font=f(15), fill=(120, 120, 120))
    im.save(path)


def banner_good(path):
    """필수 문구를 갖춘 비교적 깨끗한 광고 (오탐 확인용)."""
    im = Image.new("RGB", (1400, 900), (250, 250, 245))
    d = ImageDraw.Draw(im)
    d.text((60, 60), "든든암보험", font=f(54, True), fill=(20, 60, 120))
    d.text((60, 150), "암 진단 시 5,000만원을 지급합니다", font=f(52, True), fill=(30, 30, 30))
    d.text((60, 250), "가입연령 20~65세 · 20년납 20년만기", font=f(40), fill=(30, 30, 30))
    d.text((60, 400), "보험계약 체결 전 상품설명서와 약관을 반드시 확인하시기 바랍니다.", font=f(30), fill=(60, 60, 60))
    d.text((60, 450), "갱신 시 보험료가 인상될 수 있으며, 보장개시일부터 90일간 면책, 1년 이내 진단 시 50% 감액 지급됩니다.", font=f(24), fill=(60, 60, 60))
    d.text((60, 800), "준법감시인 심의필 제2026-0123호 (2026.10.01)", font=f(26), fill=(60, 60, 60))
    im.save(path)


def banner_savings(path):
    im = Image.new("RGB", (1400, 700), (240, 248, 240))
    d = ImageDraw.Draw(im)
    d.text((60, 60), "차곡차곡저축보험", font=f(54, True), fill=(10, 100, 50))
    d.text((60, 160), "다른 보험사보다 저렴! 전액 환급 가능", font=f(54, True), fill=(30, 30, 30))
    d.text((60, 280), "5년납 · 가입연령 0~75세 · 확정 3.5% 이자", font=f(40), fill=(30, 30, 30))
    d.text((60, 600), "※ 공시이율은 변동될 수 있습니다.", font=f(16), fill=(110, 110, 110))
    im.save(path)


def make_pptx(path, img_path):
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    # 슬라이드 1: 네이티브 텍스트 (과장/수치 오류 + 작은 글씨)
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.6), Inches(12), Inches(1.5))
    tb.text_frame.text = "차곡차곡저축보험 - 원금 보장! 확정 3.5% 이자"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    tb2 = s.shapes.add_textbox(Inches(0.6), Inches(2.4), Inches(12), Inches(1.2))
    tb2.text_frame.text = "5년납 · 가입연령 0~75세 · 누구나 가입"
    tb2.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    tb3 = s.shapes.add_textbox(Inches(0.6), Inches(6.6), Inches(12), Inches(0.5))
    tb3.text_frame.text = "※ 공시이율은 변동될 수 있으며 약관을 확인하시기 바랍니다."
    tb3.text_frame.paragraphs[0].runs[0].font.size = Pt(6)

    # 슬라이드 2: 이미지(배너)만 삽입
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(str(img_path), Inches(0.5), Inches(0.5), width=Inches(12.3))

    # 슬라이드 3: 심의필 (예금자보호/해지환급금은 일부러 없음), 그룹 도형 없이 텍스트
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.6), Inches(6.5), Inches(12), Inches(0.6))
    tb.text_frame.text = "준법감시인 심의필 제2026-0456호"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    prs.save(str(path))


def make_reference_docs():
    import docx

    d = docx.Document()
    d.add_heading("(가상) 보험광고 심의 규정", 1)
    arts = [
        ("제1조", "목적", "이 규정은 보험상품 광고가 소비자를 오인하게 하지 않도록 심의 기준을 정함을 목적으로 한다."),
        ("제2조", "심의필 표기", "광고물에는 준법감시인의 심의필 번호와 심의일자를 표기하여야 한다."),
        ("제3조", "과장 광고 금지", "광고에서 최고, 최대, 1위 등 최상급 표현이나 100%, 무조건 등 단정적 표현은 객관적 근거와 기준일을 함께 표시하지 아니하는 한 사용할 수 없다."),
        ("제4조", "보장범위 안내", "보험금 지급 사유, 면책기간, 감액지급 기간 등 보장 제한 사항은 보장내용과 동일한 크기로 명확히 알려야 한다."),
        ("제5조", "예금자보호 안내", "저축성 보험 및 연금보험 광고에는 예금자보호법에 따른 보호 여부 및 보호 한도를 안내하여야 한다."),
        ("제6조", "해지환급금 안내", "저축성 보험 광고에는 해지환급금이 납입보험료보다 적거나 없을 수 있음을 안내하여야 한다."),
        ("제7조", "타사 비교 금지", "다른 보험회사의 상품과 비교하여 우월함을 나타내는 표현을 사용할 수 없다."),
        ("제8조", "갱신형 상품 안내", "갱신형 상품 광고에는 갱신 시 보험료가 인상될 수 있음을 안내하여야 한다."),
        ("제9조", "글자 크기", "유의사항 등 소비자에게 불리한 내용은 광고 본문의 가장 큰 글자의 일정 비율 이상으로 표기하여야 한다."),
    ]
    for no, title, body in arts:
        d.add_paragraph(f"{no}({title}) {body}")
    d.save(str(R / "sample_ad_review_rules.docx"))
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(R),
                    str(R / "sample_ad_review_rules.docx")], capture_output=True, timeout=180)


if __name__ == "__main__":
    banner_bad(S / "ad_cancer_bad.png")
    banner_good(S / "ad_cancer_good.png")
    banner_savings(S / "banner_savings.png")
    make_pptx(S / "ad_savings_deck.pptx", S / "banner_savings.png")
    make_reference_docs()
    print("샘플 생성 완료:", sorted(p.name for p in list(S.iterdir()) + list(R.iterdir())))
