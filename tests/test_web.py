"""웹 화면·기준 관리 테스트 (OCR 은 가짜 엔진 사용)."""
import io
import shutil
import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from PIL import Image

from adreview import criteria as C
from adreview.models import TextBlock
from adreview.ocr import OCREngine
from adreview.webapp import State, create_app

ROOT = Path(__file__).resolve().parent.parent


class FakeEngine(OCREngine):
    name = "fake"

    def recognize(self, image_path):
        return [TextBlock(text="든든암보험 최고의 보장 무조건 100% 지급", bbox=(10, 10, 300, 40), height_px=30),
                TextBlock(text="준법감시인 심의필 제2026-1호", bbox=(10, 60, 300, 80), height_px=20)]


def make_app():
    tmp = Path(tempfile.mkdtemp())
    cfg = tmp / "config"
    shutil.copytree(ROOT / "config", cfg, ignore=shutil.ignore_patterns("history"))
    st = State("paddle", {}, config_dir=cfg, out_root=tmp / "out", engine_factory=lambda n: FakeEngine())
    return create_app(st).test_client(), st, tmp


def png_bytes():
    b = io.BytesIO()
    Image.new("RGB", (400, 120), "white").save(b, "PNG")
    return b.getvalue()


def wait_done(c, jid):
    for _ in range(100):
        j = c.get(f"/job/{jid}/status").get_json()
        if j["state"] in ("done", "failed"):
            return j
        time.sleep(0.1)
    raise AssertionError("검수가 끝나지 않음")


def test_pages_render():
    c, st, tmp = make_app()
    for url in ("/", "/jobs", "/criteria", "/criteria?t=checklist", "/criteria?t=master", "/criteria?t=history"):
        r = c.get(url)
        assert r.status_code == 200, url
    assert "버전" in c.get("/").get_data(as_text=True)


def test_upload_check_shows_criteria_in_report():
    c, st, tmp = make_app()
    r = c.post("/check", data={"files": (io.BytesIO(png_bytes()), "광고.png"), "engine": "paddle", "today": "2026-10-16"},
               content_type="multipart/form-data")
    jid = r.get_json()["url"].rsplit("/", 1)[1]
    j = wait_done(c, jid)
    assert j["state"] == "done", j
    html = c.get(f"/job/{jid}/file/report.html").get_data(as_text=True)
    assert "이번 검수에 적용된 기준" in html and "광고.png" in html and "BAN-001" in html
    crit = c.get(f"/job/{jid}/file/criteria.json").get_json()
    assert crit["version"] == C.describe(C.load("rules", st.config_dir), C.load("checklist", st.config_dir),
                                         C.load("master", st.config_dir), config_dir=st.config_dir)["version"]
    assert not list((st.out_root / jid / "uploads").glob("*"))      # 원본은 보관하지 않음


def test_rejects_bad_extension():
    c, st, tmp = make_app()
    r = c.post("/check", data={"files": (io.BytesIO(b"x"), "a.exe")}, content_type="multipart/form-data")
    assert "error" in r.get_json()


def test_edit_rule_changes_result_and_version():
    c, st, tmp = make_app()

    def run():
        r = c.post("/check", data={"files": (io.BytesIO(png_bytes()), "a.png")}, content_type="multipart/form-data")
        jid = r.get_json()["url"].rsplit("/", 1)[1]
        wait_done(c, jid)
        return jid, c.get(f"/job/{jid}/file/result.json").get_json()

    jid1, res1 = run()
    ids1 = {f["rule_id"] for f in res1[0]["findings"] if f["status"] == "review"}
    assert "BAN-002" in ids1
    v1 = c.get(f"/job/{jid1}/file/criteria.json").get_json()["version"]

    rules = C.load("rules", st.config_dir)
    ban = rules["prohibited_expressions"]
    for r in ban:
        if r["id"] == "BAN-002":
            r["enabled"] = False
    j = c.post("/api/criteria/rules/structured", json={"required_phrases": rules["required_phrases"],
                                                        "prohibited_expressions": ban, "note": "BAN-002 끔", "user": "tester"}).get_json()
    assert j["ok"], j
    jid2, res2 = run()
    ids2 = {f["rule_id"] for f in res2[0]["findings"] if f["status"] == "review"}
    assert "BAN-002" not in ids2 and "BAN-001" in ids2
    v2 = c.get(f"/job/{jid2}/file/criteria.json").get_json()["version"]
    assert v1 != v2
    assert C.list_history(config_dir=st.config_dir)[0]["note"] == "BAN-002 끔"


def test_invalid_rule_is_rejected_and_file_untouched():
    c, st, tmp = make_app()
    before = C.read_text("rules", st.config_dir)
    j = c.post("/api/criteria/rules/structured", json={
        "required_phrases": [], "prohibited_expressions": [{"id": "BAN-X", "pattern": "(깨진", "severity": "high"}]}).get_json()
    assert not j["ok"] and "정규식" in j["error"]
    assert C.read_text("rules", st.config_dir) == before
    j = c.post("/api/criteria/master/text", json={"text": "products: [ {"}).get_json()
    assert not j["ok"] and "YAML" in j["error"]


def test_checklist_disable_item_and_restore():
    c, st, tmp = make_app()
    first = C.load("checklist", st.config_dir)["items"][0]["num"]
    j = c.post("/api/criteria/checklist/apply", json={"enabled": {str(first): False}, "note": "항목 끔"}).get_json()
    assert j["ok"], j
    assert C.describe({}, C.load("checklist", st.config_dir), {}, config_dir=st.config_dir)["counts"]["disabled"] == 1
    assert c.get(f"/api/criteria/checklist/item/{first}").status_code == 200
    hist = C.list_history("checklist", config_dir=st.config_dir)
    assert c.get(f"/criteria/diff?name={hist[0]['backup']}").status_code == 200
    j = c.post("/api/criteria/restore", json={"name": hist[0]["backup"]}).get_json()
    assert j["ok"]
    assert C.describe({}, C.load("checklist", st.config_dir), {}, config_dir=st.config_dir)["counts"]["disabled"] == 0


def test_backup_name_traversal_blocked():
    c, st, tmp = make_app()
    assert c.get("/criteria/diff?name=../../etc/passwd").status_code == 404
    assert not c.post("/api/criteria/restore", json={"name": "../rules.yaml"}).get_json()["ok"]
    assert c.get("/job/..%2f..%2fetc/file/passwd").status_code == 404


# ---------------- 로그인 · 권한 · 감사 로그 ----------------
def make_auth_app(pw="user-pw", admin="admin-pw"):
    tmp = Path(tempfile.mkdtemp())
    cfg = tmp / "config"
    shutil.copytree(ROOT / "config", cfg, ignore=shutil.ignore_patterns("history"))
    st = State("paddle", {}, config_dir=cfg, out_root=tmp / "out", engine_factory=lambda n: FakeEngine(),
               password=pw, admin_password=admin)
    return create_app(st).test_client(), st


def test_login_required_and_roles():
    c, st = make_auth_app()
    assert c.get("/").status_code == 302 and "/login" in c.get("/").headers["Location"]
    assert c.get("/job/x/status").status_code == 401
    assert c.post("/api/criteria/rules/structured", json={}).status_code == 401
    # 잘못된 비밀번호
    assert "올바르지" in c.post("/login", data={"name": "홍길동", "pw": "x"}).get_data(as_text=True)
    # 검수자: 검수 화면은 되지만 기준 수정은 불가
    assert c.post("/login", data={"name": "홍길동", "pw": "user-pw"}).status_code == 302
    assert c.get("/").status_code == 200 and c.get("/criteria").status_code == 200
    r = c.post("/api/criteria/rules/structured", json={"required_phrases": [], "prohibited_expressions": []})
    assert r.status_code == 403
    # 관리자: 수정 가능, 변경자는 로그인 이름으로 기록
    c.get("/logout")
    c.post("/login", data={"name": "관리자A", "pw": "admin-pw"})
    rules = C.load("rules", st.config_dir)
    j = c.post("/api/criteria/rules/structured", json={"required_phrases": rules["required_phrases"],
               "prohibited_expressions": rules["prohibited_expressions"], "note": "t", "user": "위조"}).get_json()
    assert j["ok"] and C.list_history(config_dir=st.config_dir)[0]["user"] == "관리자A"
    log = (st.out_root / "audit.log").read_text(encoding="utf-8")
    assert "login_fail" in log and "criteria_save" in log and "관리자A" in log


def test_login_throttle():
    c, st = make_auth_app()
    for _ in range(5):
        c.post("/login", data={"name": "a", "pw": "bad"})
    r = c.post("/login", data={"name": "a", "pw": "user-pw"})       # 올바른 비밀번호여도 잠금 중
    assert r.status_code == 200 and "초과" in r.get_data(as_text=True)


def test_single_password_allows_everything():
    c, st = make_auth_app(pw="only", admin="")
    c.post("/login", data={"name": "김", "pw": "only"})
    rules = C.load("rules", st.config_dir)
    assert c.post("/api/criteria/rules/structured", json={"required_phrases": rules["required_phrases"],
                  "prohibited_expressions": rules["prohibited_expressions"]}).get_json()["ok"]


def test_default_ocr_opts_use_bundled_models():
    from adreview.ocr import default_ocr_opts
    o = default_ocr_opts()
    assert Path(o["det_model_dir"]).name == "PP-OCRv5_mobile_det" and Path(o["rec_model_dir"]).exists()


# ---------------- 보안 ----------------
def test_filename_is_escaped_no_stored_xss():
    c, st, tmp = make_app()
    evil = "<img src=x onerror=alert(1)>.png"
    r = c.post("/check", data={"files": (io.BytesIO(png_bytes()), evil), "user": "<script>x</script>"},
               content_type="multipart/form-data")
    jid = r.get_json()["url"].rsplit("/", 1)[1]
    wait_done(c, jid)
    for url in ("/jobs", f"/job/{jid}"):
        h = c.get(url).get_data(as_text=True)
        assert "<img src=x" not in h and "<script>x</script>" not in h, url
    rep = c.get(f"/job/{jid}/file/report.html")
    assert "<img src=x" not in rep.get_data(as_text=True)
    assert "default-src 'none'" in rep.headers["Content-Security-Policy"]


def test_security_headers():
    c, st, tmp = make_app()
    h = c.get("/").headers
    assert h["X-Content-Type-Options"] == "nosniff" and h["Cache-Control"] == "no-store"
    assert "frame-ancestors" in h["Content-Security-Policy"] and h["X-Frame-Options"] == "SAMEORIGIN"


def test_redos_pattern_rejected():
    c, st, tmp = make_app()
    j = c.post("/api/criteria/rules/structured", json={"required_phrases": [], "prohibited_expressions": [
        {"id": "BAN-X", "pattern": "(a+)+$", "severity": "high"}]}).get_json()
    assert not j["ok"] and "중첩" in j["error"]
    assert not C.validate_rules(C.load("rules", ROOT / "config"))


def test_xlsx_formula_injection_neutralized():
    from openpyxl import load_workbook
    from adreview.models import AdPage
    from adreview.checks import check_file
    from adreview.report import write_xlsx
    pg = AdPage(file="=HYPERLINK(\"http://evil\")", label="image",
                blocks=[TextBlock(text="=cmd|' /C calc'!A0 최고 무조건", source="native")])
    res = check_file([pg], C.load("rules", ROOT / "config"), {})
    out = Path(tempfile.mkdtemp())
    wb = load_workbook(write_xlsx([res], str(out)))
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    assert cell.data_type == "s", (ws.title, cell.coordinate)


def test_input_guards():
    import zipfile
    from adreview.extract import guard_pptx, extract_any
    from adreview.ocr import MAX_PIXELS, load_bgr
    tmp = Path(tempfile.mkdtemp())
    z = tmp / "bomb.pptx"
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as f:
        f.writestr("ppt/big.bin", b"\0" * (10 * 1024 * 1024))
    guard_pptx(str(z))                                     # 10MB 는 통과
    import adreview.extract as E
    old, E.MAX_PPTX_UNZIPPED = E.MAX_PPTX_UNZIPPED, 1024 * 1024
    try:
        try:
            guard_pptx(str(z))
            raise AssertionError("압축 폭탄이 통과됨")
        except ValueError:
            pass
    finally:
        E.MAX_PPTX_UNZIPPED = old
    import adreview.ocr as O
    big = tmp / "big.png"
    Image.new("L", (100, 100)).save(big)
    old, O.MAX_PIXELS = O.MAX_PIXELS, 50
    try:
        try:
            O.load_bgr(str(big))
            raise AssertionError("큰 이미지가 통과됨")
        except ValueError:
            pass
    finally:
        O.MAX_PIXELS = old


def test_retention_purge():
    import os
    c, st, tmp = make_app()
    old = st.out_root / "20200101-000000-abcdef"
    old.mkdir()
    os.utime(old, (1, 1))
    new = st.out_root / "20990101-000000-abcdef"
    new.mkdir()
    st.retention_days = 30
    assert st.purge_old() == 1 and not old.exists() and new.exists()


if __name__ == "__main__":
    fails = 0
    for n, fn in list(globals().items()):
        if n.startswith("test_") and callable(fn):
            try:
                fn()
                print("PASS", n)
            except Exception as e:
                fails += 1
                import traceback
                traceback.print_exc()
                print("FAIL", n, "->", repr(e))
    sys.exit(1 if fails else 0)
